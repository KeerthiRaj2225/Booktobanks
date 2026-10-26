"""Excel and CSV exports."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..parser.models import Statement

COLUMNS = [
    ("Date", 12), ("Value Date", 12), ("Narration", 60), ("Ref / Chq No", 20), ("Party", 26),
    ("Debit", 14), ("Credit", 14), ("Balance", 16), ("Ledger", 24), ("Voucher Type", 13), ("Check", 10),
]
INR = '#,##0.00;[Red]-#,##0.00'


def _row(t):
    return [
        t.date, t.value_date, t.narration, t.ref, t.party,
        t.debit or None, t.credit or None, t.balance, t.ledger, t.voucher_type,
        "OK" if t.reconciled else ("Review" if t.reconciled is False else ""),
    ]


def to_csv(st: Statement) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c for c, _ in COLUMNS])
    for t in st.transactions:
        r = _row(t)
        r[0] = t.date.strftime("%d-%m-%Y")
        r[1] = t.value_date.strftime("%d-%m-%Y") if t.value_date else ""
        r[5] = f"{t.debit:.2f}" if t.debit else ""
        r[6] = f"{t.credit:.2f}" if t.credit else ""
        r[7] = f"{t.balance:.2f}" if t.balance is not None else ""
        w.writerow(r)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens ₹/Unicode correctly


def to_excel(st: Statement, bank_ledger: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")
    review_fill = PatternFill("solid", fgColor="FDECEC")
    thin = Side(style="thin", color="D9DEE5")

    for i, (name, width) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22

    for r, t in enumerate(st.transactions, 2):
        for ci, v in enumerate(_row(t), 1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.border = Border(bottom=thin)
            if ci in (1, 2) and v:
                cell.number_format = "DD-MM-YYYY"
            if ci in (6, 7, 8):
                cell.number_format = INR
            if ci == 3:
                cell.alignment = Alignment(wrap_text=False)
        if t.reconciled is False:
            for ci in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=ci).fill = review_fill

    last = len(st.transactions) + 1
    tr = last + 1
    ws.cell(row=tr, column=5, value="Total").font = Font(bold=True)
    for col in (6, 7):
        L = get_column_letter(col)
        c = ws.cell(row=tr, column=col, value=f"=SUM({L}2:{L}{last})")
        c.number_format, c.font = INR, Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    s = wb.create_sheet("Summary")
    rows = [
        ("Bank", st.bank),
        ("Account number", st.account_number),
        ("IFSC", st.ifsc),
        ("Account holder", st.account_holder),
        ("Tally bank ledger", bank_ledger),
        ("Period", f"{st.period_from or ''} to {st.period_to or ''}".strip()),
        ("Opening balance", st.opening_balance),
        ("Total debits", f"=Transactions!F{tr}"),
        ("Total credits", f"=Transactions!G{tr}"),
        ("Closing balance", st.closing_balance),
        ("Computed closing", "=B7-B8+B9"),
        ("Transactions", len(st.transactions)),
        ("Rows matching running balance", f"{st.reconciled_pct}%"),
    ]
    for i, (k, v) in enumerate(rows, 1):
        s.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = s.cell(row=i, column=2, value=v)
        if k in {"Opening balance", "Total debits", "Total credits", "Closing balance", "Computed closing"}:
            c.number_format = INR
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 34
    if st.warnings:
        s.cell(row=len(rows) + 2, column=1, value="Warnings").font = Font(bold=True, color="B42318")
        for j, w in enumerate(st.warnings):
            s.cell(row=len(rows) + 3 + j, column=1, value=w)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
