import io
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.main import app  # noqa: E402
from app.parser.utils import parse_amount, parse_date  # noqa: E402
from tests.make_samples import build_all  # noqa: E402

SAMPLES = ROOT / "samples"
client = TestClient(app)


@pytest.fixture(scope="session", autouse=True)
def samples():
    build_all()


def _convert(name, password=None):
    pdf = (SAMPLES / name).read_bytes()
    data = {"password": password} if password else {}
    return client.post("/api/convert", files={"file": (name, pdf, "application/pdf")}, data=data)


@pytest.mark.parametrize("name", ["hdfc.pdf", "sbi.pdf", "icici.pdf", "axis.pdf", "kotak.pdf", "kotak_amount_desc.pdf"])
def test_every_bank_matches_truth(name):
    r = _convert(name)
    assert r.status_code == 200, r.text
    body = r.json()
    truth = json.loads((SAMPLES / name.replace(".pdf", ".truth.json")).read_text())
    assert body["count"] == len(truth["transactions"])
    assert body["reconciled_pct"] == 100.0
    assert abs(body["opening_balance"] - truth["opening"]) < 0.01
    assert abs(body["closing_balance"] - truth["closing"]) < 0.01
    for got, exp in zip(body["transactions"], truth["transactions"]):
        assert got["date"] == exp["date"]
        assert abs(got["debit"] - exp["debit"]) < 0.01
        assert abs(got["credit"] - exp["credit"]) < 0.01


def test_password_flow():
    assert _convert("hdfc_locked.pdf").json()["error"] == "password_required"
    assert _convert("hdfc_locked.pdf", "wrong").json()["error"] == "wrong_password"
    ok = _convert("hdfc_locked.pdf", "DEMO1234")
    assert ok.status_code == 200 and ok.json()["count"] == 25


def test_rejects_non_pdf():
    r = client.post("/api/convert", files={"file": ("x.pdf", b"hello", "application/pdf")})
    assert r.status_code == 415


def test_exports():
    st = _convert("hdfc.pdf").json()
    payload = {"statement": st, "bank_ledger": "HDFC Bank A/c", "company": "Demo Co"}

    x = client.post("/api/export/xlsx", json=payload)
    wb = load_workbook(io.BytesIO(x.content))
    assert wb["Transactions"].max_row == st["count"] + 2  # header + rows + total

    c = client.post("/api/export/csv", json=payload)
    assert c.content.decode("utf-8-sig").count("\n") == st["count"] + 1

    v = client.post("/api/export/tally-vouchers", json=payload)
    root = ET.fromstring(v.content)
    vouchers = root.findall(".//VOUCHER")
    assert len(vouchers) == st["count"]
    for vch in vouchers:  # every voucher must balance to zero
        total = sum(float(a.text) for a in vch.findall(".//AMOUNT"))
        assert abs(total) < 0.001

    m = client.post("/api/export/tally-masters", json=payload)
    names = [l.get("NAME") for l in ET.fromstring(m.content).findall(".//LEDGER")]
    assert "HDFC Bank A/c" in names and "Cash" not in names


def test_utils():
    assert parse_amount("1,23,456.78") == 123456.78
    assert parse_amount("500.00 Dr") == -500.0
    assert parse_date("1 Apr 2024").isoformat() == "2024-04-01"
    assert parse_date("01-Apr-24").isoformat() == "2024-04-01"
    assert parse_date("31/03/25").isoformat() == "2025-03-31"
